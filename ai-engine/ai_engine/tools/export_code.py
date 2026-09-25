"""Export a readable copy of the source code for written submission.

Walks the project (React frontend + Python AI engine) and writes
``submission/CODE_SUBMISSION.txt`` and ``submission/CODE_SUBMISSION.docx``,
skipping generated folders and binary files. The DOCX uses a monospace font
for the code, a heading per file, page breaks and a simple file index.
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

# Root of the whole project (contains ai-engine and Save-the-planet-main).
PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent.parent
SUBMISSION_DIR = PROJECT_ROOT / 'submission'

SKIP_DIRS = {'node_modules', '.git', 'venv', '.venv', 'artifacts', 'dist',
             'build', '__pycache__', '.pytest_cache', 'data', 'reports',
             'submission', '.idea', '.next'}
BINARY_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.ico', '.mp3',
                     '.mp4', '.wav', '.woff', '.woff2', '.ttf', '.otf',
                     '.pyc', '.joblib', '.docx', '.pkl', '.webp', '.svg',
                     '.lock'}
TEXT_EXTENSIONS = {'.py', '.ts', '.tsx', '.css', '.js', '.json', '.md', '.html',
                   '.txt', '.toml', '.yaml', '.yml', '.ini', '.cfg', '.sh',
                   '.bat'}


def collect_files():
    """Return a sorted, filtered list of candidate source files."""
    files = []
    if not PROJECT_ROOT.exists():
        return files
    for path in PROJECT_ROOT.rglob('*'):
        if not path.is_file():
            continue
        rel = path.relative_to(PROJECT_ROOT)
        if any(part in SKIP_DIRS for part in rel.parts):
            continue
        if path.suffix.lower() in BINARY_EXTENSIONS:
            continue
        if path.suffix.lower() not in TEXT_EXTENSIONS:
            continue
        files.append(path)
    return sorted(files)


def _try_read_text(path):
    try:
        return path.read_text(encoding='utf-8', errors='replace')
    except Exception:
        return ''


def write_txt(files):
    lines = ['ECOQUEST SOURCE EXPORT', '=' * 40, '',
             f'Total files: {len(files)}', '']
    for i, path in enumerate(files, 1):
        rel = path.relative_to(PROJECT_ROOT)
        lines.append('')
        lines.append(f'File {i}: {rel}')
        lines.append('-' * 40)
        lines.append(_try_read_text(path))
        lines.append('')
    SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
    (SUBMISSION_DIR / 'CODE_SUBMISSION.txt').write_text('\n'.join(lines), encoding='utf-8')


def write_docx(files):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Consolas'
    normal.font.size = Pt(9)

    doc.add_heading('EcoQuest: Learn & Save the Planet - Source Code', level=0)
    doc.add_paragraph('File index (in order of appearance):')
    for i, path in enumerate(files, 1):
        rel = path.relative_to(PROJECT_ROOT)
        doc.add_paragraph(f'{i}. {rel}', style='List Number')

    for path in files:
        rel = path.relative_to(PROJECT_ROOT)
        doc.add_heading(str(rel), level=2)
        code = _try_read_text(path)
        for line in code.splitlines():
            para = doc.add_paragraph(line)
            for run in para.runs:
                run.font.name = 'Consolas'
                run.font.size = Pt(8)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(SUBMISSION_DIR / 'CODE_SUBMISSION.docx')


def main():
    files = collect_files()
    print(f'Exporting {len(files)} source files...')
    write_txt(files)
    write_docx(files)
    print(f'Wrote {SUBMISSION_DIR / "CODE_SUBMISSION.txt"}')
    print(f'Wrote {SUBMISSION_DIR / "CODE_SUBMISSION.docx"}')


if __name__ == '__main__':
    sys.exit(main())